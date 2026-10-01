import openpyxl as xl

#loading the sheet and getting the cell
wb = xl.load_workbook('transactions.xlsx')
sheet = wb['Sheet1']
cell = sheet.cell(1,1)
print(sheet.max_row)
for row in range(1,sheet.max_row+1):
    print(sheet.cell(row,1).value)
